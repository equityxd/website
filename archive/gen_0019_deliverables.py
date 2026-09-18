# -*- coding: utf-8 -*-
"""
Regenerates serial CV-20260914-0019 deliverables (5 cover letters + 6 interview-prep docs)
for the target role 'Director of Finance' (Indeed aggregation JD).

All content is built strictly from facts present in Ernest SONG's base CV.
No experience, metric, company or location is invented.
"""

import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SERIAL = "CV-20260914-0019"
BASE = os.path.dirname(os.path.abspath(__file__))
CL_DIR = os.path.join(BASE, "cover_letters")
IP_DIR = os.path.join(BASE, "6 Interview Prep")


def add_title(doc, lines):
    """lines = (text, size, bold, (r,g,b))"""
    for i, (text, size, bold, color) in enumerate(lines):
        p = doc.add_paragraph()
        p.alignment = [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
                         WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT][i]
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.font.bold = bold
        if color is not None:
            r.font.color.rgb = RGBColor(*color)


def para(doc, text, size=10, bold=False, italic=False, color=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color is not None:
        r.font.color.rgb = RGBColor(*color)
    return p


def hr(doc):
    p = doc.add_paragraph()
    r = p.add_run("─" * 78)
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
    return p


# ---------------------------------------------------------------------------
# COVER LETTERS (5). Each bridges the candidate's real experience to the
# JD's top requirements and includes an honest, non-apologetic gap statement.
# All bodies kept under 250 words.
# ---------------------------------------------------------------------------

COVER_LETTERS = {
    "CL1": {
        "subject": "Director of Finance — financial reporting, close & multi-entity GL",
        "body": [
            "Dear Hiring Manager,",
            "",
            "I am applying for the Director of Finance role. The position's need for reliable monthly and quarterly financial reporting, disciplined budgeting & forecasting, and P&L-adjacent decision-making across multiple entities matches my 15+ years across BE, FR and NL.",
            "",
            "As Business Analyst on the SAP S/4HANA (Engie SEM) and SAP HANA (Holcim) migrations, I owned finance sub-process design and reconciled multi-entity data into the new systems. I automate month-end close with PowerQuery ETL (1,500+ bookings per close, zero manual intervention) and turn raw data into executive Power BI and Cognos reporting. At Rexel I led IFRS financial reporting and budgeting & forecasting for Asia-Pacific, Latin American and Canadian subsidiaries.",
            "",
            "I am honest about two gaps: explicit day-to-day AP/AR bookkeeping as a standing role, and US GAAP/SEC context, which I have not used directly. I cover the first with my ESCP auditing qualification and fast-learning profile, and the second with a structured onboarding ramp — my IFRS and BE-GAAP consolidation foundation transfers cleanly. I respect deadlines, communicate clearly from the board down, and work well on-site or remotely.",
            "",
            "I would welcome a conversation to discuss how I can contribute.",
            "",
            "Best regards,",
            "Ernest SONG",
        ],
    },
    "CL2": {
        "subject": "Director of Finance — budgeting, forecasting & strategic finance",
        "body": [
            "Dear Hiring Manager,",
            "",
            "I am applying for the Director of Finance position. The JD's emphasis on budgeting & forecasting, strategic planning and cross-functional partnering mirrors my career: I repeatedly build finance planning functions from scratch and translate them into executive decision-support.",
            "",
            "At Engie SEM I delivered a budgeting & forecasting model from scratch within a two-week window, replacing an unreliable legacy framework to protect continuity. at Vinci Airports I built long-term business models covering macroeconomic impact, CapEx and concession valuation, and negotiated contract extensions that improved budgeting and forecasting. At Rexel I led annual budgeting, monthly forecasting and variance analysis versus budgeted results across subsidiaries.",
            "",
            "I also own strategic planning and P&L-adjacent partnership, having sat on both the operator's and the executive-recruiter's side of the table. I bring calm under pressure, rigorous Power BI reporting, and a record of turning finance into a lever for growth. I am transparent about two smaller gaps — a US GAAP/SEC cadence and the explicit finance-team management structure — and address them with a clear ramp plan and a proven ability to become productive quickly. I am available immediately and flexible on location.",
            "",
            "I would welcome the opportunity to discuss my planning and partnering experience.",
            "",
            "Best regards,",
            "Ernest SONG",
        ],
    },
    "CL3": {
        "subject": "Director of Finance — P&L partnering, cash flow & leadership",
        "body": [
            "Dear Hiring Manager,",
            "",
            "I am applying for the Director of Finance role, drawn by its blend of P&L responsibility, cash-flow oversight and team leadership.",
            "",
            "As interim CFO and fundraising consultant at Magnetrap I operated under direct P&L and treasury oversight: I arranged €3M in debt and equity financing, prepared cash-flow projections and business plans, and implemented corrective actions to protect profitability. at Engie Tractebel I designed cash-flow reporting and per-project cash exposure, performed counterpart credit analysis and impairment testing, and built finance data models from SAP HANA into Power BI reporting.",
            "",
            "I have also led and mentored cross-functional teams — audit teams at KPMG and operations across 50+ airport sites at Vinci Airports — partnering comfortably with non-finance stakeholders. I will be honest about the gaps I have not held in a US setting: day-to-day AP/AR bookkeeping as a standing function and US GAAP/SEC reporting. My ESCP auditing qualification, repeated success building finance functions from scratch, and autonomous, fast-learning style give me a concrete plan to close those quickly. I communicate clearly from senior management to end users and respect all deadlines.",
            "",
            "I would welcome a conversation to discuss how I can contribute to your finance function.",
            "",
            "Best regards,",
            "Ernest SONG",
        ],
    },
    "CL4": {
        "subject": "Director of Finance — data-to-reporting & cross-functional partnering",
        "body": [
            "Dear Hiring Manager,",
            "",
            "I am applying for the Director of Finance role, where reliable financial reporting, data-driven budgeting and cross-functional partnering are essential.",
            "",
            "My key differentiator is turning raw finance data into executive decision-support: I have rolled out Power BI, Cognos, QlikSense and Business Objects reporting on top of SAP HANA and SAP BPC across multiple entities. at Holcim I integrated Qliksense to turn SAP data into actionable reporting, and at Engie SEM I bridge IT architecture and business operations as the core Business Analyst during the SAP S/4HANA migration.",
            "",
            "I pair this with operational finance ownership — month-end close automation, budgeting & forecasting, cash-flow and treasury oversight, and multi-entity GL — so I speak fluently to both finance and the wider business. I am straightforward about the areas I have not held in a US context: daily AP/AR bookkeeping and US GAAP/SEC reporting. I cover these with my formal auditing qualification, a proven record of building finance and reporting functions from scratch, and a structured onboarding ramp. I work well on-site or remotely, respect deadlines, and communicate clearly at all levels.",
            "",
            "I would welcome a conversation to discuss how I can contribute.",
            "",
            "Best regards,",
            "Ernest SONG",
        ],
    },
    "CL5": {
        "subject": "Director of Finance — operator-to-executive perspective",
        "body": [
            "Dear Hiring Manager,",
            "",
            "I am applying for the Director of Finance role. What sets me apart is perspective: I have sat on both the operator's and the executive-recruiter's side of the table, so I read the numbers and the organisation behind them.",
            "",
            "Across 15+ years I have led financial reporting, month-end close leadership and budgeting & forecasting for multi-entity, multi-country operations. I have managed SAP S/4HANA and SAP HANA migrations, automated close with PowerQuery ETL (1,500+ bookings per close, zero manual intervention), and built executive Power BI and Cognos reporting. I have also operated with P&L and treasury oversight as interim CFO, raising €3M and protecting profitability through corrective action.",
            "",
            "I will be honest about two gaps for a US role: explicit day-to-day AP/AR bookkeeping as a standing function and US GAAP/SEC reporting. I cover the first with my ESCP auditing qualification and fast-learning profile, and the second through a structured onboarding ramp — my IFRS/BE-GAAP consolidation foundation transfers cleanly. I respect deadlines, communicate clearly from the board to end users, and am available immediately with flexible interview availability.",
            "",
            "I would welcome a conversation to discuss how I can contribute to your finance function.",
            "",
            "Best regards,",
            "Ernest SONG",
        ],
    },
}

# ---------------------------------------------------------------------------
# INTERVIEW PREP (6). STAR answers grounded strictly in the base CV.
# ---------------------------------------------------------------------------

INTERVIEW_QUESTIONS = {
    "IP1": "Tell me about a time you led finance-process integration during an ERP change.",
    "IP2": "Give an example of how you automate a repetitive financial process.",
    "IP3": "How do you make raw data usable for senior decision-makers?",
    "IP4": "Describe how you own budgeting & forecasting and variance analysis.",
    "IP5": "Give an example of leading multi-entity consolidation / financial reporting.",
    "IP6": "Tell me about a time you had to build something from scratch under a tight deadline.",
}

STAR = {
    "IP1": {
        "Situation": "Engie SEM — SAP S/4HANA migration pilot across French subsidiaries.",
        "Task": "Ensure finance sub-processes were modelled and implemented correctly in the new ERP.",
        "Action": "Acted as the core Business Analyst bridging IT architecture and business operations; drove AS-IS to TO-BE process design; ran finance workshops; reconciled the middle-office (IVDB) data with S/4HANA and simplified WDS project coding.",
        "Result": "A clean finance transition with a reconciled, monitorable project structure — directly reusable for any ERP replacement rollout.",
    },
    "IP2": {
        "Situation": "Monthly financial close was ingesting 1,500+ bookings, done manually.",
        "Task": "Remove manual effort while keeping the close accurate and auditable.",
        "Action": "Built a PowerQuery ETL pipeline to ingest bookings automatically and reconcile the middle-office (IVDB) data.",
        "Result": "1,500+ bookings per close processed with zero manual intervention, shortening the close and reducing error risk.",
    },
    "IP3": {
        "Situation": "Raw finance data that executives needed to act on quickly.",
        "Task": "Turn raw data into clear executive reporting.",
        "Action": "Built Power BI dashboards and rolled out Qliksense and Cognos finance reporting on top of SAP HANA and SAP BPC.",
        "Result": "Faster, clearer decision-support for leadership across multiple entities.",
    },
    "IP4": {
        "Situation": "Need for reliable budgeting, forecasting and performance tracking.",
        "Task": "Own the planning cycle and track performance against plan.",
        "Action": "At Rexel led annual budgeting, monthly forecasting and variance analysis versus budgeted results; at Engie SEM built a budgeting & forecasting model from scratch within a two-week window, replacing an unreliable legacy framework.",
        "Result": "More reliable planning and early detection of deviations, protecting operational continuity.",
    },
    "IP5": {
        "Situation": "A listed, multi-entity group needed consistent consolidated financial reporting.",
        "Task": "Lead IFRS financial reporting across geographies.",
        "Action": "At Rexel led IFRS reporting for Asia-Pacific, Latin America and Canadian subsidiaries and integrated SAP BPC and Cognos reporting systems to strengthen consolidation.",
        "Result": "Consolidated, audit-ready multi-entity reporting and a stronger consolidation platform.",
    },
    "IP6": {
        "Situation": "A critical resource gap threatened budget continuity at Engie SEM.",
        "Task": "Deliver a robust budgeting & forecasting tool within a two-week timeline.",
        "Action": "Designed and built the budget model from scratch using PowerQuery, replacing an unreliable legacy framework, and streamlined financial-close automation.",
        "Result": "Operational continuity protected and a repeatable, automated planning tool delivered on an aggressive deadline.",
    },
}


def write_cover_letter(num):
    key = "CL" + num
    data = COVER_LETTERS[key]
    doc = Document()
    std = doc.styles["Normal"]
    std.font.size = Pt(11)
    std.font.name = "Source Sans 3"

    add_title(doc, [
        ("COVER LETTER", 16, True, (0, 0, 0)),
        ("Serial CV-20260914-0019", 12, False, (0x55, 0x55, 0x55)),
        ("─" * 78, 8, False, (0xAA, 0xAA, 0xAA)),
        (data["subject"], 11, True, (0x33, 0x33, 0x33)),
    ])
    for line in data["body"]:
        bold = line.startswith("Best regards,") or line == "Ernest SONG"
        para(doc, line, 10, bold=bold)

    path = os.path.join(CL_DIR, f"{SERIAL}_CL{num}.docx")
    doc.save(path)
    return path


def write_interview_prep(num):
    key = "IP" + num
    star = STAR[key]
    doc = Document()
    std = doc.styles["Normal"]
    std.font.size = Pt(11)
    std.font.name = "Source Sans 3"

    add_title(doc, [
        ("INTERVIEW PREP (STAR)", 16, True, (0, 0, 0)),
        ("Serial CV-20260914-0019", 12, False, (0x55, 0x55, 0x55)),
        ("─" * 78, 8, False, (0xAA, 0xAA, 0xAA)),
    ])
    para(doc, "INTERVIEW PREP — Serial CV-20260914-0019 (STAR)", 11, True, color=(0x33, 0x33, 0x33))
    para(doc, "Role: Director of Finance (US) — JD: Indeed 'Director of Finance' aggregation.", 9, False, italic=True, color=(0x77, 0x77, 0x77))
    hr(doc)
    para(doc, INTERVIEW_QUESTIONS[key], 11, True)
    for heading in ["Situation", "Task", "Action", "Result"]:
        para(doc, f"{heading.upper()}:", 11, True, color=(0x33, 0x33, 0x33))
        para(doc, star[heading], 10, False, space_after=4)

    path = os.path.join(IP_DIR, f"{SERIAL}_IP{num}.docx")
    doc.save(path)
    return path


def main():
    os.makedirs(CL_DIR, exist_ok=True)
    os.makedirs(IP_DIR, exist_ok=True)
    written = []
    for num in ["1", "2", "3", "4", "5"]:
        p = write_cover_letter(num)
        written.append(p)
        print("WROTE", p)
    for num in ["1", "2", "3", "4", "5", "6"]:
        p = write_interview_prep(num)
        written.append(p)
        print("WROTE", p)
    print(f"\nTotal files written: {len(written)}")


if __name__ == "__main__":
    main()
